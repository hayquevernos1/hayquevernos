"""Catálogo v2 (Diego, 14 oct 2026) → catalogo_v2.js
Fuente: Catalogo_eventos_concentrado_v2.xlsx (Matriz maestra + Catering alimentos/bebidas + Venues).
Homologado: nombres corregidos, Equipamiento repartido como «Solo renta» dentro de cada categoría,
huecos llenados (postres, pirotecnia, inflables, photo booth, limpieza, oficiante), cada categoría etiquetada S/E."""
import json, re, sys, unicodedata
from pathlib import Path
import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else '/tmp/xl3/c.xlsx'
OUT = Path('/home/claude/hqv/v5/catalogo_v2.js')

L = lambda s: [x.strip() for x in s.split(',') if x.strip()]
# n, emoji, alcance (S social, E empresarial), servicios completos, staff (plural), solo renta
CATS = [
 ("Catering", "🍽️", "SE",
  "Banquete, Buffet, Cóctel, Estación, Box lunch, Barra de bebidas, Coffee break, Mesa de postres y pastel, Comida para staff y proveedores",
  "Chefs, Cocineros, Parrilleros, Taqueros, Reposteros, Meseros, Garroteros, Capitanes de meseros, Bartenders, Mixólogos, Baristas, Sommeliers, Ayudantes de cocina",
  "Equipo de cocina, Chafers y calentadores, Barras móviles, Dispensadores y vitroleros"),
 ("Música", "🎵", "SE",
  "DJ, Solistas, Dúos y tríos, Ensambles, Bandas, Mariachis, Orquestas, Coros, Grupos versátiles",
  "DJ, Cantantes, Instrumentistas, Músicos de sesión, Directores musicales, Coristas, Percusionistas",
  "Instrumentos musicales"),
 ("Entretenimiento y shows", "🎭", "SE",
  "Espectáculos artísticos, Magia, Comedia, Danza, Teatro, Shows infantiles, Shows temáticos, Espectáculos interactivos, Pirotecnia y efectos especiales, Inflables y juegos infantiles",
  "Animadores, Magos, Comediantes, Bailarines, Coreógrafos, Actores, Mimos, Acróbatas, Malabaristas, Recreadores, Monitores infantiles",
  "Inflables y brincolines, Juegos infantiles, Máquinas de efectos (humo, burbujas, confeti)"),
 ("Infraestructura", "🏗️", "SE",
  "Carpas, Escenarios, Tarimas, Estructuras, Sanitarios, Sistemas eléctricos temporales, Climatización, Instalaciones provisionales",
  "Montadores, Instaladores, Electricistas, Técnicos de estructuras, Tramoyistas, Técnicos de rigging, Supervisores de montaje",
  "Carpas, Tarimas y pistas, Sanitarios portátiles, Plantas de luz, Climas y calentadores, Herramientas"),
 ("Mobiliario y montaje", "🪑", "SE",
  "Montaje de mobiliario, Mobiliario temático, Mobiliario corporativo",
  "Montadores de mobiliario, Cargadores, Auxiliares de montaje, Acomodadores, Supervisores de montaje",
  "Mesas, Sillas, Salas lounge, Barras, Mostradores, Mobiliario temático, Mantelería"),
 ("Cristalería y cubertería", "🍷", "SE",
  "Montaje de mesa completo",
  "Personal de montaje de mesas, Auxiliares de servicio, Lavaloza, Stewards",
  "Vajillas, Cristalería, Cubiertos, Menaje, Utensilios, Accesorios de mesa"),
 ("Audio, iluminación, foto y video", "📸", "SE",
  "Sonorización, Iluminación, Fotografía social, Fotografía corporativa, Videografía, Producción audiovisual, Streaming, Video mapping, Photo booth y cabina 360, Servicios multimedia",
  "Ingenieros de audio, Técnicos de sonido, Iluminadores, Fotógrafos, Camarógrafos, Operadores de cámara, Microfonistas, Editores audiovisuales, Operadores de streaming",
  "Bocinas, Micrófonos, Consolas, Iluminación, Cámaras, Pantallas y proyectores"),
 ("Tecnología", "💻", "E",
  "Internet temporal, Conectividad y redes, Sistemas de registro, Software para eventos, Experiencias digitales, Realidad virtual, Realidad aumentada, Automatización, Soporte tecnológico",
  "Técnicos de redes, Técnicos de soporte TI, Operadores de sistemas, Técnicos de hardware, Operadores de experiencias digitales",
  "Computadoras, Tablets, Equipo de red"),
 ("Decoración y ambientación", "🌸", "SE",
  "Decoración floral, Ambientación, Escenografía, Decoración temática, Instalaciones artísticas, Decoración con globos, Branding de espacios",
  "Decoradores, Floristas, Diseñadores florales, Escenógrafos, Instaladores de decoración, Globoflexistas",
  "Centros de mesa, Backdrops y muros, Arcos y bases"),
 ("Marketing y comunicación", "📣", "E",
  "Campañas, Branding, Activaciones de marca, Relaciones públicas, Publicidad, Gestión de influencers, Contenido digital, Comunicación corporativa",
  "Promotores, Demostradores, Embajadores de marca, Community managers, Creadores de contenido, Encuestadores, Asesores comerciales",
  ""),
 ("Transporte", "🚐", "SE",
  "Traslado de invitados, Transporte ejecutivo, Transporte de grupos, Transporte de carga, Shuttles, Vehículos especiales, Valet parking",
  "Choferes, Conductores ejecutivos, Operadores de transporte, Valet parking, Auxiliares de estacionamiento, Coordinadores de transporte",
  "Autos clásicos y especiales"),
 ("Hospedaje", "🏨", "SE",
  "Alojamiento individual, Alojamiento grupal, Hoteles, Resorts, Hospedaje ejecutivo, Bloques de habitaciones, Alojamiento para artistas",
  "Recepcionistas, Concierges, Botones, Personal de hospitalidad, Asistentes de huéspedes",
  ""),
 ("Maquillaje y estilismo", "💄", "SE",
  "Maquillaje social, Maquillaje nupcial, Maquillaje editorial y artístico, Peinado, Barbería, Manicure, Servicios de belleza, Asesoría de imagen",
  "Maquillistas, Estilistas, Peinadores, Barberos, Manicuristas, Asesores de imagen, Vestuaristas",
  ""),
 ("Registro y protocolo", "🎟️", "SE",
  "Registro de asistentes, Acreditación, RSVP, Control de accesos, Recepción, Atención VIP, Protocolo social, Protocolo corporativo, Maestro de ceremonias, Oficiante de ceremonias",
  "Hostess, Edecanes, Recepcionistas, Acreditadores, Maestros de ceremonias, Moderadores, Auxiliares de registro, Coordinadores de protocolo",
  ""),
 ("Exposición y negocios", "🏢", "E",
  "Stands, Pabellones, Showrooms, Exhibiciones comerciales, Matchmaking empresarial, Gestión de expositores, Networking",
  "Asesores comerciales, Demostradores, Capturistas de leads, Ejecutivos de ventas, Asistentes de stand, Facilitadores de networking",
  "Stands modulares"),
 ("Souvenirs y regalos", "🎁", "SE",
  "Recuerdos, Regalos corporativos, Artículos promocionales, Kits, Productos personalizados, Empaques, Personalización en sitio",
  "Diseñadores, Ilustradores, Personalizadores, Empacadores, Auxiliares de producción, Personal de entrega",
  ""),
 ("Seguridad, salud y protección civil", "🛡️", "SE",
  "Seguridad privada, Vigilancia, Protección ejecutiva, Cobertura médica, Ambulancias, Brigadas, Planes de emergencia, Protección civil",
  "Guardias de seguridad, Escoltas, Supervisores de seguridad, Paramédicos, Enfermeros, Médicos, Brigadistas, Rescatistas, Salvavidas, Especialistas en protección civil",
  ""),
 ("Actividades y experiencias", "🎯", "SE",
  "Talleres, Catas y degustaciones, Clases, Actividades deportivas, Experiencias culturales, Experiencias gastronómicas, Actividades de integración, Team building, Bienestar",
  "Instructores, Talleristas, Facilitadores, Guías, Coaches de actividades, Monitores recreativos, Entrenadores, Especialistas en experiencias",
  ""),
 ("Impresión y producción gráfica", "🖨️", "SE",
  "Invitaciones, Impresión digital, Gran formato, Señalética, Rotulación, Materiales publicitarios, Acreditaciones impresas, Displays",
  "Diseñadores gráficos, Impresores, Rotulistas, Instaladores de vinil, Operadores de impresión, Auxiliares de producción gráfica",
  ""),
 ("Vestuario y accesorios", "👗", "SE",
  "Confección de vestidos, Confección de trajes, Venta de vestidos y trajes, Disfraces, Uniformes, Vestuario escénico, Accesorios y complementos",
  "Sastres, Modistas, Costureros, Vestuaristas, Asesores de vestuario, Asistentes de vestuario",
  "Vestidos, Trajes, Disfraces, Accesorios"),
 ("Permisos y trámites", "📋", "E",
  "Gestión de permisos, Licencias, Seguros para eventos, Contratos, Derechos musicales, Trámites, Asesoría jurídica, Cumplimiento normativo",
  "Abogados, Gestores de permisos, Asesores administrativos, Especialistas en cumplimiento, Auxiliares administrativos",
  ""),
 ("Sustentabilidad y limpieza", "♻️", "SE",
  "Limpieza del evento, Gestión de residuos, Reciclaje, Compostaje, Eventos cero residuos, Medición de huella ambiental, Consultoría ambiental",
  "Personal de limpieza, Consultores ambientales, Supervisores de residuos, Clasificadores de materiales, Auxiliares de reciclaje",
  ""),
]

def fix_g(g):
    g = (g or '').strip()
    m = {'asiatica': 'Asiática', 'medio oriente': 'Medio Oriente', 'mediterranea': 'Mediterránea',
         'mar y tierra': 'Mar y tierra', 'mixta': 'Fusión e internacional'}
    return m.get(g.lower(), g[:1].upper() + g[1:])

def clean_items(s):
    s = re.sub(r'\.\s*Los refrescos.*$', '', str(s or '')).strip().rstrip('.')
    return [x[:1].upper() + x[1:] for x in L(s) if x.lower() != 'otro' and 'no aplica' not in x.lower()]

wb = openpyxl.load_workbook(SRC, read_only=True)
# ---- cocina (una sola vez; las 5 pestañas comparten el mismo catálogo)
coc = {'Salado': {}, 'Dulce': {}}
for r in wb['Catering - Alimentos - Banquete'].iter_rows(min_row=3, values_only=True):
    if not r[0]: continue
    tipo, g, esp, pl = r[3], fix_g(r[5]), (r[6] or '').strip(), r[7]
    if not esp or esp.lower() == 'otro': continue
    coc[tipo].setdefault(g, {})[esp] = clean_items(pl)
# ---- bebidas
beb = {}
for r in wb['Catering - Bebidas'].iter_rows(min_row=3, values_only=True):
    if not r[0]: continue
    alc, temp, tipo, var = r[2], r[3], r[5], clean_items(r[6])
    if not var: continue
    t = beb.setdefault(tipo, {'n': tipo, 'alc': [], 'v': []})
    if alc not in t['alc']: t['alc'].append(alc)
    t['v'] += [x for x in var if x not in t['v']]
# ---- venues
venues = [r[1] for r in wb['Venues'].iter_rows(min_row=2, values_only=True) if r and r[1]]

data = {
    'cats': [{'n': n, 'e': e, 't': t, 'full': L(f), 'staff': L(s), 'rent': L(rt)} for n, e, t, f, s, rt in CATS],
    'cocina': coc,
    'bebidas': list(beb.values()),
    'bebidasServicio': ['Barra', 'Estación', 'Mesa', 'Carrito', 'Dispensador', 'Barril'],
    'venues': venues,
}
OUT.write_text('/* Catálogo v2 ¡Hay que vernos! — generado por gen_cat2.py */\nwindow.HQV_CAT2=' + json.dumps(data, ensure_ascii=False, separators=(',', ':')) + ';\n')
ns = sum(len(v) for g in coc.values() for v in g.values())
print('ok', OUT.stat().st_size, 'bytes ·', len(data['cats']), 'categorías ·', ns, 'especialidades ·', len(beb), 'tipos de bebida ·', venues)
